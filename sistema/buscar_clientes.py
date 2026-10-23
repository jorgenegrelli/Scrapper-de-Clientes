"""
Busca empresas no Google Maps (via o motor local gosom/google-maps-scraper) e monta uma planilha Excel
so com quem NAO tem site proprio: sem site, ou que usa Instagram / WhatsApp / Facebook / Linktree /
site gratuito no lugar de um site. Cada linha ja vem com link do WhatsApp, perfil do Google e avaliacoes.

Uso:
    python buscar_clientes.py --nicho "dentista" --nicho "clinica odontologica" --cidade "Curitiba, PR" --minimo 80
    python buscar_clientes.py --modelo        (so recria a planilha MODELO em branco)

O motor precisa estar ligado (sistema/iniciar-motor.ps1; o BUSCAR CLIENTES.bat ja liga sozinho).
"""
import argparse, csv, datetime, json, math, os, re, sys, urllib.parse
from html import escape

PASTA = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # raiz do projeto
PLANILHAS = os.path.join(PASTA, "Planilhas")
BRUTOS = os.path.join(PASTA, "sistema", "dados-brutos")
MODELO = os.path.join(PASTA, "MODELO - Planilha de Clientes.xlsx")
MENSAGENS = os.path.join(PASTA, "mensagens")
CONFIG = os.path.join(PASTA, "configuracao.txt")

# As mensagens ficam em mensagens/*.txt (edite la). Campos trocados automaticamente:
# {seu_nome} (de configuracao.txt) {empresa} {nota} {avaliacoes} {avaliacoes_aprox} {nicho} {cidade} {complemento}
# {saudacao}: "Bom dia"/"Boa tarde"/"Boa noite" pela hora do clique no painel ("Olá" na planilha).
# mensagem-destaque.txt vale para quem tem nota >= 4.5 e pelo menos 20 avaliacoes; o resto usa mensagem-padrao.txt.
COMPLEMENTO = {  # trecho que entra em {complemento}, conforme a coluna "Situacao do site"
    "So Instagram": ", só o Instagram",
    "So Facebook": ", só a página no Facebook",
    "So WhatsApp": ", só o link do WhatsApp",
    "Link de bio (Linktree etc.)": ", só um link de bio",
    "Site gratuito / subdominio": ", só uma página simples",
    "So perfil em plataforma": ", só o perfil numa plataforma de agendamento",
}


def ler_texto(caminho):
    try:
        with open(caminho, encoding="utf-8-sig") as f:
            return f.read().strip()
    except OSError:
        return ""


def ler_config():
    cfg = {}
    for linha in ler_texto(CONFIG).splitlines():
        k, sep, v = linha.partition("=")
        if sep and not k.strip().startswith("#"):
            cfg[k.strip()] = v.strip()
    return cfg


SEU_NOME = ler_config().get("seu_nome", "")
MENSAGEM_DESTAQUE = ler_texto(os.path.join(MENSAGENS, "mensagem-destaque.txt"))
MENSAGEM_PADRAO = ler_texto(os.path.join(MENSAGENS, "mensagem-padrao.txt"))

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import motor_api as scrape  # noqa: E402  (geocode, grid_cells, run_job, dedupe, in_bbox)

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

# Dominios que NAO contam como site proprio -> rotulo que vai para a coluna "Situacao do site"
SEM_SITE_PROPRIO = [
    (("instagram.com", "instagr.am"), "So Instagram"),
    (("wa.me", "whatsapp.com", "wa.link", "whats.link"), "So WhatsApp"),
    (("facebook.com", "fb.com", "fb.me"), "So Facebook"),
    (("linktr.ee", "linkbio", "bio.link", "beacons.ai", "lnk.bio", "taplink", "linkin.bio",
      "campsite.bio", "linkme.bio", "zaap.bio"), "Link de bio (Linktree etc.)"),
    (("business.site", "sites.google.com", "g.page", "negocio.site", "wixsite.com", "webnode",
      "weebly.com", "blogspot.", "wordpress.com", "carrd.co", "my.canva.site", "godaddysites.com"),
     "Site gratuito / subdominio"),
    (("doctoralia.", "boaconsulta.", "agendarconsulta.", "ifood.", "booksy.", "tuasaude."),
     "So perfil em plataforma"),
    (("google.com", "goo.gl", "maps.app"), "So link do Google"),
]

STATUS = ["Novo", "Contatado", "Respondeu", "Reuniao marcada", "Proposta enviada", "Fechado", "Sem interesse"]

COLUNAS = [  # (titulo, largura)
    ("Status", 16), ("Empresa", 34), ("Categoria", 22), ("Situacao do site", 24), ("Link atual", 30),
    ("Telefone", 17), ("WhatsApp", 16), ("Tipo de numero", 24), ("Nota", 7), ("Avaliacoes", 11),
    ("Perfil no Google", 16), ("Mensagem", 60), ("Ultimas avaliacoes", 60), ("Bairro", 18),
    ("Endereco", 40), ("Observacoes", 30),
]


def situacao_site(url):
    """'' se tem site proprio; senao o rotulo do que a empresa usa no lugar."""
    url = (url or "").strip()
    if not url:
        return "Sem site"
    host = urllib.parse.urlparse(url if "://" in url else "http://" + url).netloc.lower()
    for dominios, rotulo in SEM_SITE_PROPRIO:
        if any(d in host for d in dominios):
            return rotulo
    return ""


def whatsapp(telefone, site):
    """-> (link, tipo). Prefere o numero do link de WhatsApp que a empresa colocou no Google."""
    m = re.search(r"(?:wa\.me/|phone=)(\d{10,13})", site or "")
    if m:
        d, tipo = m.group(1), "WhatsApp informado no Google"
    else:
        d = re.sub(r"\D", "", telefone or "")
        if d.startswith("55") and len(d) in (12, 13):
            d = d[2:]
        d = d.lstrip("0")
        if len(d) == 10 and d[2] in "6789":  # celular antigo sem o 9
            d = d[:2] + "9" + d[2:]
        if len(d) == 11 and d[2] == "9":
            tipo = "Celular (WhatsApp provavel)"
        elif len(d) == 10:
            tipo = "Fixo (testar WhatsApp Business)"
        else:
            return "", "Sem numero valido"
        d = "55" + d
    link = f"https://web.whatsapp.com/send?phone={d}"  # abre direto no WhatsApp Web (wa.me passa por uma tela antes)
    return link, tipo


def ultimas_avaliacoes(raw, n=3):
    try:
        revs = json.loads(raw or "null") or []
    except ValueError:
        return ""
    revs.sort(key=lambda r: r.get("published_at") or "", reverse=True)
    linhas = []
    for r in revs[:n]:
        texto = (r.get("text_original") or r.get("Description") or "").strip().replace("\n", " ")
        if len(texto) > 180:
            texto = texto[:177] + "..."
        quando = (r.get("published_at") or "")[:10] or r.get("When", "")
        linhas.append(f"{r.get('Rating', '?')}★ {r.get('Name', '')} ({quando}): {texto or '(sem texto)'}")
    return "\n".join(linhas)


def bairro(raw):
    try:
        return (json.loads(raw or "{}") or {}).get("borough", "")
    except ValueError:
        return ""


def perfil_google(r):
    if r.get("place_id"):
        q = urllib.parse.quote(r.get("title", ""))
        return f"https://www.google.com/maps/search/?api=1&query={q}&query_place_id={r['place_id']}"
    return r.get("link", "")


def nome_curto(titulo):
    """'Dra. Fulana | Dentista em Curitiba | Aparelho' -> 'Dra. Fulana' (fica natural na mensagem)."""
    return re.split(r"\s+[|–-]\s+", titulo or "")[0].strip() or titulo


class Campos(dict):
    """Campo desconhecido no texto da mensagem fica como esta, em vez de dar erro."""
    def __missing__(self, k):
        return "{" + k + "}"


def aproximado(n):
    """471 -> 'mais de 450', como uma pessoa falaria. Numeros pequenos ficam exatos."""
    passo = 10 if n < 100 else 50 if n < 1000 else 100
    base = n // passo * passo
    if n < 20 or base == n:
        return f"{n:,}".replace(",", ".")
    return "mais de " + f"{base:,}".replace(",", ".")


def mensagem(empresa, sit, nota, qtd, nicho, cidade):
    """{saudacao} continua no texto: o painel troca pela hora do clique, a planilha por 'Ola'."""
    modelo = MENSAGEM_DESTAQUE if nota and nota >= 4.5 and qtd >= 20 else MENSAGEM_PADRAO
    if not modelo:
        return ""
    return modelo.format_map(Campos(
        seu_nome=SEU_NOME, empresa=nome_curto(empresa), nota=f"{nota:.1f}".replace(".", ",") if nota else "",
        avaliacoes=qtd, avaliacoes_aprox=aproximado(qtd), nicho=nicho, cidade=cidade,
        complemento=COMPLEMENTO.get(sit, "")))


def para_lead(r, nicho, cidade):
    sit = situacao_site(r.get("website"))
    wa, tipo = whatsapp(r.get("phone"), r.get("website"))
    try:
        nota = float(r.get("review_rating") or 0) or None
    except ValueError:
        nota = None
    try:
        qtd = int(float(r.get("review_count") or 0))
    except ValueError:
        qtd = 0
    msg_painel = mensagem(r.get("title", ""), sit, nota, qtd, nicho, cidade)
    msg = msg_painel.replace("{saudacao}", "Olá")
    wa_base = wa
    if wa and msg:
        wa += "&text=" + urllib.parse.quote(msg)
    return {
        "Status": "Novo", "Empresa": r.get("title", ""), "Categoria": r.get("category", ""),
        "Situacao do site": sit, "Link atual": r.get("website", ""), "Telefone": r.get("phone", ""),
        "WhatsApp": wa, "Tipo de numero": tipo, "Nota": nota, "Avaliacoes": qtd,
        "Perfil no Google": perfil_google(r), "Mensagem": msg, "Ultimas avaliacoes": ultimas_avaliacoes(r.get("user_reviews")),
        "Bairro": bairro(r.get("complete_address")), "Endereco": r.get("address", ""), "Observacoes": "",
        "_wa_base": wa_base, "_mensagem_painel": msg_painel,  # so para o painel, nao vao para a planilha
    }


def salvar_planilha(leads, caminho, titulo):
    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"
    cab = PatternFill("solid", fgColor="1F4E78")
    for i, (nome, larg) in enumerate(COLUNAS, 1):
        c = ws.cell(row=1, column=i, value=nome)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), cab
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = larg
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "C2"

    link = Font(color="0563C1", underline="single")
    textos_link = {"WhatsApp": "Abrir WhatsApp", "Perfil no Google": "Ver no Google", "Link atual": None}
    for lin, lead in enumerate(leads, 2):
        for col, (nome, _) in enumerate(COLUNAS, 1):
            v = lead.get(nome, "")
            c = ws.cell(row=lin, column=col)
            if nome in textos_link and v:
                c.value, c.hyperlink, c.font = textos_link[nome] or v, v, link
            else:
                c.value = v
            c.alignment = Alignment(vertical="top",
                                    wrap_text=nome in ("Mensagem", "Ultimas avaliacoes", "Endereco", "Empresa"))
        ws.cell(row=lin, column=9).number_format = "0.0"

    ultima = max(len(leads) + 1, 500)
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"A2:A{ultima}")
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(COLUNAS)).column_letter}{max(len(leads) + 1, 2)}"

    info = wb.create_sheet("Sobre")
    for k, v in [("Busca", titulo), ("Gerada em", datetime.datetime.now().strftime("%d/%m/%Y %H:%M")),
                 ("Empresas", len(leads)), ("", ""),
                 ("Situacao do site", "O que a empresa colocou no campo 'site' do Google (vazio = Sem site)."),
                 ("Tipo de numero", "Celular quase sempre tem WhatsApp. Fixo pode ter WhatsApp Business: teste o link."),
                 ("Status", "Use a lista suspensa para acompanhar o contato com cada cliente.")]:
        info.append([k, v])
    info.column_dimensions["A"].width, info.column_dimensions["B"].width = 18, 90
    for c in info["A"]:
        c.font = Font(bold=True)
    wb.save(caminho)


PAINEL_MODELO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "painel_modelo.html")


def numero_internacional(wa):
    """'...send?phone=5541995120077' -> '+55 41 99512-0077' (formato que o WhatsApp reconhece e deixa tocar)."""
    m = re.search(r"phone=55(\d{2})(\d{4,5})(\d{4})$", wa or "")
    return f"+55 {m[1]} {m[2]}-{m[3]}" if m else ""


def salvar_painel(leads, caminho, titulo):
    """Pagina local: 'Abrir conversa' copia a mensagem e reaproveita sempre a mesma aba do WhatsApp Web."""
    dados = [{"empresa": l["Empresa"], "nome": nome_curto(l["Empresa"]),
              "numero": numero_internacional(l["_wa_base"]) or l["Telefone"], "nota": f"{l['Nota']:.1f}".replace(".", ",") if l["Nota"] else "",
              "avaliacoes": l["Avaliacoes"], "situacao": l["Situacao do site"], "bairro": l["Bairro"],
              "telefone": l["Telefone"], "tipo": l["Tipo de numero"], "wa": l["_wa_base"], "id": l["_wa_base"] or l["Empresa"],
              "mensagem": l["_mensagem_painel"], "google": l["Perfil no Google"], "link": l["Link atual"],
              "avaliacoes_txt": l["Ultimas avaliacoes"]} for l in leads]
    js = json.dumps(dados, ensure_ascii=False).replace("</", "<\\/")
    with open(PAINEL_MODELO, encoding="utf-8") as f:
        modelo = f.read()
    html = (modelo.replace("__TITULO__", escape(titulo))
                  .replace("__GERADO__", datetime.datetime.now().strftime("%d/%m/%Y"))
                  .replace("__DADOS__", js))
    with open(caminho, "w", encoding="utf-8") as f:
        f.write(html)


def buscar(nichos, cidade, minimo, grid, depth, max_time):
    termos = [f"{n} em {cidade.replace(',', '')}" for n in nichos]
    print(f"▶ Localizando {cidade}...")
    geo = scrape.geocode(cidade, with_bbox=True)
    if not geo or not geo[2]:
        sys.exit("✗ Nao achei a cidade. Escreva como 'Curitiba, PR'.")
    lat, lon, bbox = geo
    cells = scrape.grid_cells(bbox, grid)
    # do centro para fora: se o minimo for atingido cedo, a cobertura fica no miolo da cidade
    cells.sort(key=lambda c: math.hypot(float(c["lat"]) - float(lat), float(c["lon"]) - float(lon)))
    try:
        scrape.req("GET", "/api/v1/jobs")
    except Exception:
        sys.exit("✗ O motor esta desligado. Rode iniciar-motor.ps1 na pasta do kit e tente de novo.")

    print(f"▶ Termos: {', '.join(termos)}")
    print(f"▶ {len(cells)} areas da cidade, uma de cada vez. Meta: {minimo} empresas sem site proprio.")
    print("  (Buscas grandes e seguidas podem fazer o Google bloquear seu IP por algumas horas.)\n")
    body = {"name": "clientes-sites", "keywords": termos, "lang": "pt", "fast_mode": False,
            "depth": depth, "email": False, "max_time": max_time}
    rows = []
    for i, cell in enumerate(cells, 1):
        print(f"▶ Area {i}/{len(cells)}")
        try:
            _, novas = scrape.run_job(dict(body, **cell))
        except Exception as e:
            print(f"\n  ⚠ area {i} pulada: {str(e).strip()}")
            continue
        rows = scrape.dedupe(rows + novas)
        rows = [r for r in rows if scrape.in_bbox(r, bbox)]
        sem_site = sum(1 for r in rows if situacao_site(r.get("website")))
        print(f"  {len(rows)} empresas no total, {sem_site} sem site proprio")
        if sem_site >= minimo:
            print(f"✓ Meta de {minimo} atingida.")
            break
    return rows


def main():
    ap = argparse.ArgumentParser(description="Busca clientes sem site proprio no Google Maps.")
    ap.add_argument("--nicho", action="append", help="tipo de negocio (pode repetir)")
    ap.add_argument("--cidade", help='ex: "Curitiba, PR"')
    ap.add_argument("--minimo", type=int, default=80)
    ap.add_argument("--grid", type=int, default=3, help="divide a cidade em N x N areas (padrao 3)")
    ap.add_argument("--depth", type=int, default=10)
    ap.add_argument("--max-time", type=int, default=900, help="limite por area, em segundos")
    ap.add_argument("--filtro-categoria", help="so mantem categorias que contenham um destes trechos, "
                                                "separados por virgula (ex: odont,dentist)")
    ap.add_argument("--de-arquivo", help="reprocessa um CSV de dados-brutos sem buscar de novo")
    ap.add_argument("--modelo", action="store_true", help="so recria a planilha MODELO em branco")
    ap.add_argument("--so-painel", action="store_true", help="gera so o painel HTML, sem nova planilha")
    a = ap.parse_args()

    os.makedirs(PLANILHAS, exist_ok=True)
    os.makedirs(BRUTOS, exist_ok=True)
    if a.modelo:
        salvar_planilha([], MODELO, "Modelo em branco")
        return print(f"✓ Modelo salvo: {MODELO}")
    if not (a.nicho and a.cidade):
        ap.error("informe --nicho e --cidade")

    nome = f"{a.nicho[0].title()} - {a.cidade.replace(',', '')} - {datetime.date.today():%Y-%m-%d}"
    if a.de_arquivo:
        with open(a.de_arquivo, encoding="utf-8") as f:
            rows = list(csv.DictReader(f))
    else:
        rows = buscar(a.nicho, a.cidade, a.minimo, a.grid, a.depth, a.max_time)
        if not rows:
            sys.exit("✗ Nenhuma empresa encontrada.")
        bruto = os.path.join(BRUTOS, nome + ".csv")
        with open(bruto, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0].keys()), extrasaction="ignore")
            w.writeheader()
            w.writerows(rows)

    total = len(rows)
    if a.filtro_categoria:
        trechos = [t.strip().lower() for t in a.filtro_categoria.split(",") if t.strip()]
        rows = [r for r in rows if any(t in (r.get("category") or "").lower() for t in trechos)]
    rows = [r for r in rows if "fechado" not in (r.get("status") or "").lower()
            and "closed" not in (r.get("status") or "").lower()]
    cidade_curta = a.cidade.split(",")[0].strip()
    leads = [para_lead(r, a.nicho[0], cidade_curta) for r in rows if situacao_site(r.get("website"))]
    # celular primeiro, depois quem tem mais avaliacoes (clinica movimentada sem site = melhor cliente)
    leads.sort(key=lambda l: (not l["Tipo de numero"].startswith(("Celular", "WhatsApp")), -l["Avaliacoes"]))

    titulo = f"{', '.join(a.nicho)} em {a.cidade}"
    painel = os.path.join(PLANILHAS, f"PAINEL - {nome}.html")
    salvar_painel(leads, painel, titulo)
    if a.so_painel:
        return print(f"✓ Painel: {painel}")

    # nunca grava por cima de uma planilha existente (ela pode ter Status/Observacoes preenchidos)
    caminho = os.path.join(PLANILHAS, nome + ".xlsx")
    n = 2
    while True:
        try:
            if os.path.exists(caminho):
                raise FileExistsError
            salvar_planilha(leads, caminho, titulo)
            break
        except (FileExistsError, PermissionError):
            caminho = os.path.join(PLANILHAS, f"{nome} ({n}).xlsx")
            n += 1
    if not os.path.exists(MODELO):
        salvar_planilha([], MODELO, "Modelo em branco")

    print(f"\n✓ {len(leads)} empresas sem site proprio (de {total} encontradas)")
    cont = {}
    for l in leads:
        cont[l["Situacao do site"]] = cont.get(l["Situacao do site"], 0) + 1
    for k, v in sorted(cont.items(), key=lambda x: -x[1]):
        print(f"    {k}: {v}")
    cel = sum(1 for l in leads if l["Tipo de numero"].startswith(("Celular", "WhatsApp")))
    print(f"    WhatsApp por celular: {cel} | fixo (testar): {len(leads) - cel}")
    print(f"  Planilha: {caminho}")
    print(f"  Painel (para enviar as mensagens): {painel}")


if __name__ == "__main__":
    main()
